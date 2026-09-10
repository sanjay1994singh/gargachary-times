"""Excel exports for authorized subscription administrators."""
from datetime import datetime
from io import BytesIO
import json

from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def subscription_excel_response(queryset, request):
    queryset = queryset.filter(payment_status='SUCCESS')
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('User subscriptions')
    sheet.freeze_panes = 'A2'
    profile_fields = ('email', 'mobile', 'address', 'city', 'district', 'state', 'pincode', 'country')
    fields = list(queryset.model._meta.concrete_fields)
    headers = ['Full name'] + [name.title() for name in profile_fields]
    headers += [str(field.verbose_name).replace('_', ' ').title() for field in fields]
    headers += ['Username', 'Plan name', 'Invoice number', 'Delivery status', 'Invoice PDF link']
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24

    def append(values, header=False, invoice_url=None):
        cells = []
        for value in values:
            if isinstance(value, datetime):
                if timezone.is_aware(value):
                    value = timezone.localtime(value)
                value = value.isoformat(sep=' ', timespec='seconds')
            elif isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            cell = WriteOnlyCell(sheet)
            if isinstance(value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub('', value)
                cell.data_type = 's'  # Never execute user input as Excel formulas.
            else:
                cell.value = value
            if header:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='417690')
            cells.append(cell)
        if invoice_url:
            cells[-1].hyperlink = invoice_url
            cells[-1].style = 'Hyperlink'
        sheet.append(cells)

    append(headers, header=True)
    count = 0
    for subscription in queryset.select_related('user', 'plan', 'invoice').order_by('pk').iterator(chunk_size=1000):
        user = subscription.user
        invoice = getattr(subscription, 'invoice', None)
        invoice_url = request.build_absolute_uri(reverse(
            'admin:subscriptions_usersubscription_invoice_pdf', args=[subscription.pk]
        )) if invoice else ''
        values = [user.full_name or user.get_full_name()]
        values += [getattr(user, name) for name in profile_fields]
        values += [getattr(subscription, field.attname) for field in fields]
        values += [user.username, str(subscription.plan),
                   getattr(invoice, 'invoice_number', ''), getattr(invoice, 'delivery_status', ''),
                   invoice_url]
        append(values, invoice_url=invoice_url)
        count += 1
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{count + 1}'
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="user-subscriptions.xlsx"'
    return response
