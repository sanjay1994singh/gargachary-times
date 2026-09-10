from io import BytesIO

from django.contrib.auth.models import Permission
from django.test import TestCase, RequestFactory
from django.urls import reverse
from openpyxl import load_workbook

from account.models import User
from .models import Invoice, SubscriptionPlan, UserSubscription
from .exports import subscription_excel_response


class SubscriptionExcelTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(username='staff', is_staff=True)
        self.user = User.objects.create_user(
            username='subscriber', full_name='संजय शर्मा', city='Jaipur',
            state='Rajasthan', mobile='0012345678', address='=1+1')
        plan = SubscriptionPlan.objects.create(name='Annual', price=500, duration=1)
        for index in range(2):
            UserSubscription.objects.create(user=self.user, plan=plan, amount=500,
                                            transaction_id=f'test-{index}', payment_status='SUCCESS')
        for status in ('CREATED', 'PENDING', 'AUTHORIZED', 'CAPTURED', 'FAILED', 'REFUNDED'):
            UserSubscription.objects.create(user=self.user, plan=plan, amount=500,
                                            transaction_id=status, payment_status=status)
        self.url = reverse('admin:subscriptions_usersubscription_export_excel')

    def test_permissions(self):
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(self.url).status_code, 403)

    def test_empty_export_and_selected_action(self):
        response = subscription_excel_response(UserSubscription.objects.none(), RequestFactory().get('/'))
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 1)
        self.staff.user_permissions.add(Permission.objects.get(
            codename='view_usersubscription', content_type__app_label='subscriptions'))
        self.client.force_login(self.staff)
        response = self.client.post(reverse('admin:subscriptions_usersubscription_changelist'), {
            'action': 'export_selected_excel', '_selected_action': [
                UserSubscription.objects.filter(payment_status='SUCCESS').first().pk,
                UserSubscription.objects.filter(payment_status='FAILED').first().pk,
            ],
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 2)

    def test_all_records_and_profile_values(self):
        self.staff.user_permissions.add(Permission.objects.get(
            codename='view_usersubscription', content_type__app_label='subscriptions'))
        self.client.force_login(self.staff)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 3)
        row = dict(zip([cell.value for cell in sheet[1]], sheet[2]))
        self.assertEqual(row['Full name'].value, 'संजय शर्मा')
        self.assertEqual(row['City'].value, 'Jaipur')
        self.assertEqual(row['State'].value, 'Rajasthan')
        self.assertEqual(row['Mobile'].value, '0012345678')
        self.assertEqual(row['Address'].value, '=1+1')
        self.assertEqual(row['Address'].data_type, 's')
        self.assertEqual(row['Invoice number'].value, None)
        self.assertIsNone(row['Invoice PDF link'].hyperlink)
        page = self.client.get(reverse('admin:subscriptions_usersubscription_changelist'))
        self.assertContains(page, 'Download all SUCCESS subscriptions (Excel)')
        status_column = [cell.value for cell in sheet[1]].index('Payment Status')
        self.assertEqual([row[status_column] for row in sheet.iter_rows(min_row=2, values_only=True)],
                         ['SUCCESS', 'SUCCESS'])

    def test_invoice_hyperlink_and_admin_pdf_permissions(self):
        subscription = UserSubscription.objects.filter(payment_status='SUCCESS').first()
        Invoice.objects.create(subscription=subscription, invoice_number='GT-EXPORT-001',
                               billing_name='Subscriber', amount=500)
        url = reverse('admin:subscriptions_usersubscription_invoice_pdf', args=[subscription.pk])
        self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.staff.user_permissions.add(Permission.objects.get(
            codename='view_usersubscription', content_type__app_label='subscriptions'))
        response = self.client.get(self.url, secure=True)
        sheet = load_workbook(BytesIO(response.content)).active
        cell = sheet.cell(2, sheet.max_column)
        self.assertEqual(cell.hyperlink.target, 'https://testserver' + url)
        pdf = self.client.get(url)
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))
